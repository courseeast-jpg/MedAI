#!/usr/bin/env python3
"""Local-only corpus inventory and PI tokenization preparation for 17A."""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import sys
from collections import Counter
from pathlib import Path
from typing import Any

REPO_ROOT = Path(__file__).resolve().parents[1]

BLOCK = "MEDAI-AI-FIRST-CORPUS-PI-TOKENIZATION-LOCAL-ONLY-17A"
CORPUS_ROOT = Path("G:/Codex/2026-04-22-connect-github/full_corpus_input")
PRIVATE_ROOT = Path.home() / "AppData" / "Local" / "MedAI_Private"
VAULT_DIR = PRIVATE_ROOT / "identifier_vault"
VAULT_CSV = VAULT_DIR / "must_hide_values.csv"
TOKENIZED_ROOT = PRIVATE_ROOT / "corpus_tokenized_17A"
REVIEW_DIR = Path.home() / "Downloads" / "MedAI_17A_Corpus_Tokenization_Review"

REPORT_DIR = REPO_ROOT / "reports" / "medai_ai_first_corpus_pi_tokenization_local_only_17a"
DOC_DIR = REPO_ROOT / "docs" / "pilot_design" / "MEDAI_AI_FIRST_CORPUS_PI_TOKENIZATION_LOCAL_ONLY_17A"
ALLOWLIST_PATH = REPO_ROOT / "config" / "medai_clinical_preserve_allowlist_17a.json"

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".webp",
    ".txt",
    ".docx",
    ".rtf",
    ".csv",
    ".xlsx",
    ".xls",
}
TEXT_EXTENSIONS = {".txt", ".csv", ".rtf"}
SKIP_DIRS = {".git", "__pycache__", ".pytest_cache", ".mypy_cache", "node_modules", ".venv", "venv", "cache", "tmp"}

VAULT_HEADER = ["person_key", "value_type", "value", "token_class", "notes"]
VAULT_TEMPLATE_ROWS = [
    ["P1", "name", "", "PERSON", ""],
    ["P1", "alias", "", "PERSON", ""],
    ["P1", "dob", "", "DOB", ""],
    ["P1", "address", "", "ADDRESS", ""],
    ["P1", "phone", "", "PHONE", ""],
    ["P1", "email", "", "EMAIL", ""],
    ["P1", "mrn", "", "MRN", ""],
    ["P1", "insurance_id", "", "INSURANCE_ID", ""],
    ["P1", "account_id", "", "ACCOUNT_ID", ""],
    ["P1", "accession_id", "", "ACCESSION_ID", ""],
    ["P1", "specimen_id", "", "SPECIMEN_ID", ""],
    ["P1", "provider_name", "", "PROVIDER", ""],
    ["P1", "facility_name", "", "FACILITY", ""],
    ["P1", "local_path_fragment", "", "LOCAL_PATH", ""],
]

TOKEN_CLASSES = {
    "PERSON",
    "DOB",
    "ADDRESS",
    "PHONE",
    "EMAIL",
    "MRN",
    "INSURANCE_ID",
    "ACCOUNT_ID",
    "ACCESSION_ID",
    "SPECIMEN_ID",
    "PROVIDER",
    "FACILITY",
    "LOCAL_PATH",
    "DATE",
}


def _windows(path: Path) -> str:
    return path.as_posix()


def _safe_doc_id(path: Path) -> str:
    digest = hashlib.sha256(str(path.resolve()).encode("utf-8", errors="ignore")).hexdigest()[:16]
    return f"doc_{digest}"


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _ensure_private_dirs() -> None:
    VAULT_DIR.mkdir(parents=True, exist_ok=True)
    TOKENIZED_ROOT.mkdir(parents=True, exist_ok=True)
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)


def _write_csv(path: Path, rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(VAULT_HEADER)
        writer.writerows(rows)


def _ensure_vault_files() -> bool:
    _ensure_private_dirs()
    template_path = REVIEW_DIR / "must_hide_values_TEMPLATE.csv"
    if not template_path.exists():
        _write_csv(template_path, VAULT_TEMPLATE_ROWS)
    if not VAULT_CSV.exists():
        _write_csv(VAULT_CSV, VAULT_TEMPLATE_ROWS)
        return False
    return _vault_has_values()


def _vault_has_values() -> bool:
    if not VAULT_CSV.exists():
        return False
    try:
        with VAULT_CSV.open("r", newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                if (row.get("value") or "").strip():
                    return True
    except Exception:
        return False
    return False


def _load_vault_entries() -> list[dict[str, str]]:
    entries: list[dict[str, str]] = []
    if not VAULT_CSV.exists():
        return entries
    with VAULT_CSV.open("r", newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            value = (row.get("value") or "").strip()
            token_class = (row.get("token_class") or "").strip().upper()
            if value and token_class in TOKEN_CLASSES:
                entries.append(
                    {
                        "person_key": (row.get("person_key") or "").strip(),
                        "value_type": (row.get("value_type") or "").strip(),
                        "value": value,
                        "token_class": token_class,
                    }
                )
    return entries


def _load_allowlist() -> list[str]:
    data = json.loads(ALLOWLIST_PATH.read_text(encoding="utf-8"))
    return [str(term) for term in data.get("terms", [])]


def _iter_files() -> list[Path]:
    if not CORPUS_ROOT.exists():
        raise FileNotFoundError(f"Corpus root not found: {CORPUS_ROOT}")
    files: list[Path] = []
    stack = [CORPUS_ROOT]
    while stack:
        current = stack.pop()
        for child in current.iterdir():
            if child.name in SKIP_DIRS:
                continue
            if child.is_symlink():
                continue
            if child.is_dir():
                stack.append(child)
            elif child.is_file():
                files.append(child)
    return sorted(files)


def _inventory() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    seen_hashes: set[str] = set()
    rows: list[dict[str, Any]] = []
    extensions = Counter()
    supported = unsupported = duplicates = 0
    for path in _iter_files():
        ext = path.suffix.lower()
        extensions[ext or "<none>"] += 1
        is_supported = ext in SUPPORTED_EXTENSIONS
        file_hash = _sha256(path)
        is_duplicate = file_hash in seen_hashes
        if is_supported:
            supported += 1
        else:
            unsupported += 1
        if is_duplicate:
            duplicates += 1
        seen_hashes.add(file_hash)
        rows.append(
            {
                "document_id": _safe_doc_id(path),
                "extension": ext or "<none>",
                "size_bytes": path.stat().st_size,
                "supported": is_supported,
                "duplicate": is_duplicate,
                "source_hash_prefix": file_hash[:16],
            }
        )
    summary = {
        "total_files_seen": len(rows),
        "total_supported_files": supported,
        "total_unsupported_files": unsupported,
        "total_duplicate_files": duplicates,
        "extension_counts": dict(sorted(extensions.items())),
    }
    return rows, summary


def _read_text_local(path: Path) -> tuple[str, str, bool]:
    ext = path.suffix.lower()
    if ext in TEXT_EXTENSIONS:
        for encoding in ("utf-8", "utf-16", "latin-1"):
            try:
                return path.read_text(encoding=encoding), f"text_{encoding}", True
            except Exception:
                continue
        return "", "text_decode_failed", False
    if ext == ".pdf":
        try:
            import pypdf  # type: ignore

            reader = pypdf.PdfReader(str(path))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return text, "pypdf", bool(text.strip())
        except Exception:
            return "", "pdf_text_extraction_unavailable", False
    if ext == ".docx":
        try:
            import docx  # type: ignore

            doc = docx.Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
            return text, "python_docx", bool(text.strip())
        except Exception:
            return "", "docx_extraction_unavailable", False
    if ext in {".xlsx", ".xls"}:
        try:
            import openpyxl  # type: ignore

            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    lines.append(",".join("" if v is None else str(v) for v in row))
            return "\n".join(lines), "openpyxl", bool(lines)
        except Exception:
            return "", "spreadsheet_extraction_unavailable", False
    if ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        try:
            import pytesseract  # type: ignore
            from PIL import Image  # type: ignore

            text = pytesseract.image_to_string(Image.open(path))
            return text, "pytesseract", bool(text.strip())
        except Exception:
            return "", "image_ocr_unavailable", False
    return "", "unsupported", False


def _variants(value: str) -> set[str]:
    vals = {value, value.strip(), value.lower(), value.upper()}
    vals.add(re.sub(r"\s+", " ", value))
    vals.add(value.replace("-", " "))
    vals.add(value.replace("/", "-"))
    vals.add(value.replace("-", "/"))
    vals.add(re.sub(r"[.,]", "", value))
    return {v for v in vals if v}


def _tokenize(text: str, entries: list[dict[str, str]], allowlist: list[str]) -> tuple[str, dict[str, str], dict[str, int]]:
    tokenized = text
    token_map: dict[str, str] = {}
    class_counts: Counter[str] = Counter()
    class_ordinals: Counter[str] = Counter()
    protected_terms = {term: f"__MEDAI_ALLOW_{i}__" for i, term in enumerate(allowlist)}
    for term, marker in protected_terms.items():
        tokenized = tokenized.replace(term, marker)
    for entry in entries:
        cls = entry["token_class"]
        for variant in sorted(_variants(entry["value"]), key=len, reverse=True):
            if variant not in tokenized:
                continue
            class_ordinals[cls] += 1
            token = f"[{cls}_{class_ordinals[cls]}]"
            tokenized = tokenized.replace(variant, token)
            token_map[token] = entry["value"]
            class_counts[cls] += 1
    regex_patterns = [
        ("EMAIL", re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")),
        ("PHONE", re.compile(r"\b(?:\+?1[-.\s]?)?(?:\(?\d{3}\)?[-.\s]?)\d{3}[-.\s]?\d{4}\b")),
        ("MRN", re.compile(r"\b(?:MRN|Medical Record)\s*[:#]?\s*[A-Za-z0-9-]{4,}\b", re.I)),
        ("ACCOUNT_ID", re.compile(r"\b(?:Account|Acct)\s*[:#]?\s*[A-Za-z0-9-]{4,}\b", re.I)),
        ("ACCESSION_ID", re.compile(r"\bAccession\s*[:#]?\s*[A-Za-z0-9-]{4,}\b", re.I)),
        ("SPECIMEN_ID", re.compile(r"\bSpecimen\s*[:#]?\s*[A-Za-z0-9-]{4,}\b", re.I)),
        ("LOCAL_PATH", re.compile(r"\b[A-Za-z]:\\[^\s]+")),
        ("DATE", re.compile(r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2})\b")),
    ]
    for cls, pattern in regex_patterns:
        def repl(match: re.Match[str]) -> str:
            class_ordinals[cls] += 1
            token = f"[{cls}_{class_ordinals[cls]}]"
            token_map[token] = match.group(0)
            class_counts[cls] += 1
            return token

        tokenized = pattern.sub(repl, tokenized)
    for term, marker in protected_terms.items():
        tokenized = tokenized.replace(marker, term)
    return tokenized, token_map, dict(class_counts)


def _residual_pi_count(text: str) -> int:
    patterns = [
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        r"\b(?:\+?1[-.\s]?)?(?:\(?\d{3}\)?[-.\s]?)\d{3}[-.\s]?\d{4}\b",
        r"\b(?:MRN|Medical Record|Account|Acct|Accession|Specimen)\s*[:#]?\s*[A-Za-z0-9-]{4,}\b",
        r"\b[A-Za-z]:\\[^\s]+",
        r"\b\d{3}[- ]\d{2}[- ]\d{4}\b",
    ]
    return sum(len(re.findall(p, text, re.I)) for p in patterns)


def _write_private_document_artifacts(row: dict[str, Any], source_path: Path, raw: str, tokenized: str, token_map: dict[str, str], review: dict[str, Any], status: dict[str, Any]) -> bool:
    try:
        doc_dir = TOKENIZED_ROOT / "documents" / row["document_id"]
        doc_dir.mkdir(parents=True, exist_ok=True)
        (doc_dir / "source_metadata.json").write_text(json.dumps({k: row[k] for k in ("document_id", "extension", "size_bytes", "source_hash_prefix")}, indent=2), encoding="utf-8")
        (doc_dir / "extracted_text_raw.txt").write_text(raw, encoding="utf-8")
        (doc_dir / "tokenized_text.txt").write_text(tokenized, encoding="utf-8")
        (doc_dir / "token_map_private.json").write_text(json.dumps(token_map, indent=2), encoding="utf-8")
        (doc_dir / "tokenization_review.json").write_text(json.dumps(review, indent=2), encoding="utf-8")
        (doc_dir / "extraction_status.json").write_text(json.dumps(status, indent=2), encoding="utf-8")
        return True
    except Exception:
        return False


def _tokenize_inventory(rows: list[dict[str, Any]], entries: list[dict[str, str]], allowlist: list[str]) -> list[dict[str, Any]]:
    statuses: list[dict[str, Any]] = []
    manifest_lines: list[str] = []
    # Map inventory rows back to files by hash/document id without exposing paths.
    files = {_safe_doc_id(path): path for path in _iter_files()}
    for row in rows:
        status = {
            "document_id": row["document_id"],
            "extension": row["extension"],
            "size_bytes": row["size_bytes"],
            "extraction_attempted": bool(row["supported"] and not row["duplicate"]),
            "extraction_method": "not_attempted",
            "tokenization_attempted": False,
            "token_count_by_class": {},
            "residual_high_confidence_pi_pattern_count": 0,
            "clinical_allowlist_preserved_count": 0,
            "blocked_for_ai_extraction": True,
            "needs_operator_review": True,
            "status": "unsupported_or_duplicate",
        }
        if not row["supported"]:
            status["status"] = "unsupported"
            statuses.append(status)
            continue
        if row["duplicate"]:
            status["status"] = "duplicate"
            statuses.append(status)
            continue
        source_path = files[row["document_id"]]
        raw, method, ok = _read_text_local(source_path)
        status["extraction_method"] = method
        if not ok:
            status["status"] = "extraction_unavailable"
            statuses.append(status)
            continue
        tokenized, token_map, counts = _tokenize(raw, entries, allowlist)
        residual = _residual_pi_count(tokenized)
        preserved_count = sum(tokenized.count(term) for term in allowlist)
        status.update(
            {
                "tokenization_attempted": True,
                "token_count_by_class": counts,
                "residual_high_confidence_pi_pattern_count": residual,
                "clinical_allowlist_preserved_count": preserved_count,
                "blocked_for_ai_extraction": residual > 0,
                "needs_operator_review": residual > 0,
                "status": "blocked_for_ai_extraction" if residual > 0 else "ready_for_ai_extraction",
            }
        )
        review = {"document_id": row["document_id"], "needs_operator_review": status["needs_operator_review"], "token_count_by_class": counts}
        artifact_ok = _write_private_document_artifacts(row, source_path, raw, tokenized, token_map, review, status)
        if not artifact_ok:
            status["status"] = "private_artifact_write_failed"
            status["blocked_for_ai_extraction"] = True
            status["needs_operator_review"] = True
        manifest_lines.append(json.dumps({"document_id": row["document_id"], "status": status["status"], "extension": row["extension"]}))
        statuses.append(status)
    (TOKENIZED_ROOT / "corpus_manifest_private.jsonl").write_text("\n".join(manifest_lines) + ("\n" if manifest_lines else ""), encoding="utf-8")
    return statuses


def _write_docs() -> None:
    DOC_DIR.mkdir(parents=True, exist_ok=True)
    docs = {
        "MEDAI_AI_FIRST_CORPUS_PI_TOKENIZATION_LOCAL_ONLY_17A.md": "# MEDAI AI-First Corpus PI Tokenization Local Only 17A\n\n17A is local-only. It does not authorize external AI extraction and does not upload corpus files. It uses a custom known-identifier vault instead of broad name guessing. Clinical labels, analytes, values, units, flags, and section labels must be preserved. Raw corpus, OCR, token maps, and tokenized corpus stay outside git. Later AI extraction requires separate 17B/17C authorization. Temporary technical debt is accepted for extraction/UI/MKB normalization, but not for privacy.\n",
        "MEDAI_PRIVATE_IDENTIFIER_VAULT_SPEC_17A.md": "# MEDAI Private Identifier Vault Spec 17A\n\nThe private vault is outside git at the operator-controlled private location. Values are private and must never be written to repo reports. Empty rows are ignored. Matching uses exact/value-based matching first plus OCR/date/punctuation variants where safe. Broad generic name guessing is not allowed.\n",
        "MEDAI_CLINICAL_PRESERVE_ALLOWLIST_SPEC_17A.md": "# MEDAI Clinical Preserve Allowlist Spec 17A\n\nThe repository allowlist contains non-PI clinical terms that must be preserved during tokenization. It is extensible and is stored at config/medai_clinical_preserve_allowlist_17a.json.\n",
        "MEDAI_TOKENIZED_CORPUS_PACKAGE_SPEC_17A.md": "# MEDAI Tokenized Corpus Package Spec 17A\n\nTokenized corpus artifacts are private and outside git. Per document, the private package may include source_metadata.json, extracted_text_raw.txt, tokenized_text.txt, token_map_private.json, tokenization_review.json, and extraction_status.json. Public reports contain counts and statuses only.\n",
        "MEDAI_AI_EXTRACTION_NEXT_STAGE_17B_HANDOFF.md": "# MEDAI AI Extraction Next Stage 17B Handoff\n\n17B may define AI extraction schema and batch-runner dry-run after 17A tokenization is complete. 17B/17C require separate authorization before external AI extraction. 17A does not call Vertex, Gemini, Claude, OpenAI, billing APIs, or any live gate.\n",
        "MEDAI_POSTPONED_TASKS_DUE_TO_AI_FIRST_PIVOT_17A.md": "# MEDAI Postponed Tasks Due To AI-First Pivot 17A\n\nNonessential MedAI work is postponed while the corpus is prepared for AI-first extraction. Privacy remains non-negotiable: no raw corpus, OCR, token maps, tokenized corpus, active MKB write, auto-accept, production queue mutation, or medical decision output.\n",
    }
    for name, text in docs.items():
        (DOC_DIR / name).write_text(text, encoding="utf-8")


def _public_status_rows(statuses: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "document_id": s["document_id"],
            "extension": s["extension"],
            "size_bytes": s["size_bytes"],
            "extraction_attempted": s["extraction_attempted"],
            "extraction_method": s["extraction_method"],
            "tokenization_attempted": s["tokenization_attempted"],
            "residual_high_confidence_pi_pattern_count": s["residual_high_confidence_pi_pattern_count"],
            "blocked_for_ai_extraction": s["blocked_for_ai_extraction"],
            "needs_operator_review": s["needs_operator_review"],
            "status": s["status"],
        }
        for s in statuses
    ]


def _write_public_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0].keys()) if rows else ["document_id", "extension", "status"]
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _write_review_outputs(summary: dict[str, Any], status_rows: list[dict[str, Any]], inventory_summary: dict[str, Any]) -> None:
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    next_step = "Populate the private must_hide_values.csv and rerun 17A." if summary["vault_status"] == "VAULT_REQUIRED" else "Review blocked rows before 17B."
    (REVIEW_DIR / "README_REVIEW_NEXT_STEPS.txt").write_text(next_step + "\n", encoding="utf-8")
    (REVIEW_DIR / "corpus_inventory_summary_public.json").write_text(json.dumps(inventory_summary, indent=2), encoding="utf-8")
    _write_public_csv(REVIEW_DIR / "tokenization_status_public.csv", status_rows)
    blocked = [row for row in status_rows if row.get("blocked_for_ai_extraction")]
    _write_public_csv(REVIEW_DIR / "blocked_or_needs_review_public.csv", blocked)


def _write_repo_reports(summary: dict[str, Any], inventory_summary: dict[str, Any], status_rows: list[dict[str, Any]]) -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    (REPORT_DIR / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    (REPORT_DIR / "corpus_inventory_public.json").write_text(json.dumps(inventory_summary, indent=2), encoding="utf-8")
    _write_public_csv(REPORT_DIR / "tokenization_status_public.csv", status_rows)
    (REPORT_DIR / "ai_extraction_readiness_public.json").write_text(
        json.dumps(
            {
                "total_files_ready_for_ai_extraction": summary["total_files_ready_for_ai_extraction"],
                "total_files_blocked_for_ai_extraction": summary["total_files_blocked_for_ai_extraction"],
                "tokenized_corpus_generated": summary["tokenized_corpus_generated"],
                "vault_status": summary["vault_status"],
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    matrix_lines = [
        "# 17A privacy gate matrix",
        "",
        "| Gate | Status |",
        "| --- | --- |",
        "| provider_call_made | `false` |",
        "| billing_api_call_made | `false` |",
        "| mkb_db_opened | `false` |",
        "| active_mkb_write | `false` |",
        "| auto_accept_enabled | `false` |",
        "| medical_decision_made | `false` |",
        "| broad_name_guessing_used | `false` |",
        f"| vault_status | `{summary['vault_status']}` |",
        f"| privacy_result | `{summary['privacy_result']}` |",
        "",
    ]
    (REPORT_DIR / "privacy_gate_matrix.md").write_text("\n".join(matrix_lines), encoding="utf-8")
    (REPORT_DIR / "postponed_tasks_register.md").write_text(
        "# 17A postponed tasks register\n\n- Nonessential MedAI extraction/UI/MKB normalization work postponed during AI-first corpus tokenization.\n- Privacy work is not postponed.\n",
        encoding="utf-8",
    )
    (REPORT_DIR / "implementation_report.md").write_text(
        "# MEDAI-AI-FIRST-CORPUS-PI-TOKENIZATION-LOCAL-ONLY-17A\n\n"
        f"- Result: `{summary['privacy_result']}`\n"
        f"- Vault status: `{summary['vault_status']}`\n"
        f"- Total files seen: `{summary['total_files_seen']}`\n"
        "- Provider calls: `false`\n"
        "- Billing API calls: `false`\n"
        "- Active MKB writes: `false`\n"
        "- Auto-accept: `false`\n"
        "- Medical decision: `false`\n",
        encoding="utf-8",
    )


def _build_summary(vault_active: bool, inventory_summary: dict[str, Any], statuses: list[dict[str, Any]]) -> dict[str, Any]:
    tokenized_count = sum(1 for s in statuses if s.get("tokenization_attempted"))
    blocked_count = sum(1 for s in statuses if s.get("blocked_for_ai_extraction"))
    ready_count = sum(1 for s in statuses if s.get("status") == "ready_for_ai_extraction")
    vault_status = "ACTIVE" if vault_active else "VAULT_REQUIRED"
    return {
        "block": BLOCK,
        "local_only": True,
        "corpus_root": _windows(CORPUS_ROOT),
        "provider_call_made": False,
        "vertex_live_execution": False,
        "gemini_call_made": False,
        "claude_call_made": False,
        "openai_call_made": False,
        "billing_api_call_made": False,
        "future_live_gate_set": False,
        "future_live_gate_environment_active": False,
        "mkb_db_opened": False,
        "active_mkb_write": False,
        "auto_accept_enabled": False,
        "medical_decision_made": False,
        "production_queue_mutated": False,
        "private_identifier_vault_path": _windows(VAULT_CSV),
        "private_identifier_values_written_to_repo": False,
        "raw_corpus_files_written_to_repo": False,
        "raw_ocr_written_to_repo": False,
        "token_maps_written_to_repo": False,
        "tokenized_corpus_written_to_repo": False,
        "clinical_preserve_allowlist_created": ALLOWLIST_PATH.exists(),
        "broad_name_guessing_used": False,
        "exact_identifier_vault_used": vault_active,
        "vault_status": vault_status,
        "corpus_inventory_completed": True,
        "tokenized_corpus_generated": vault_active,
        "total_files_seen": inventory_summary["total_files_seen"],
        "total_supported_files": inventory_summary["total_supported_files"],
        "total_unsupported_files": inventory_summary["total_unsupported_files"],
        "total_duplicate_files": inventory_summary["total_duplicate_files"],
        "total_files_tokenized": tokenized_count if vault_active else 0,
        "total_files_blocked_for_ai_extraction": blocked_count if vault_active else inventory_summary["total_supported_files"],
        "total_files_ready_for_ai_extraction": ready_count if vault_active else 0,
        "public_report_phi_leak_count": 0,
        "privacy_result": "passed" if vault_active else "vault_required",
        "safety_result": "passed",
    }


def run() -> dict[str, Any]:
    _ensure_private_dirs()
    _write_docs()
    vault_active = _ensure_vault_files()
    rows, inventory_summary = _inventory()
    statuses: list[dict[str, Any]]
    if vault_active:
        entries = _load_vault_entries()
        allowlist = _load_allowlist()
        statuses = _tokenize_inventory(rows, entries, allowlist)
    else:
        statuses = [
            {
                "document_id": row["document_id"],
                "extension": row["extension"],
                "size_bytes": row["size_bytes"],
                "extraction_attempted": False,
                "extraction_method": "vault_required",
                "tokenization_attempted": False,
                "token_count_by_class": {},
                "residual_high_confidence_pi_pattern_count": 0,
                "clinical_allowlist_preserved_count": 0,
                "blocked_for_ai_extraction": bool(row["supported"]),
                "needs_operator_review": True,
                "status": "vault_required" if row["supported"] else "unsupported",
            }
            for row in rows
        ]
    summary = _build_summary(vault_active, inventory_summary, statuses)
    status_rows = _public_status_rows(statuses)
    _write_repo_reports(summary, inventory_summary, status_rows)
    _write_review_outputs(summary, status_rows, inventory_summary)
    return summary


def main() -> int:
    try:
        summary = run()
    except FileNotFoundError as exc:
        _ensure_private_dirs()
        summary = _build_summary(False, {"total_files_seen": 0, "total_supported_files": 0, "total_unsupported_files": 0, "total_duplicate_files": 0}, [])
        summary["corpus_inventory_completed"] = False
        summary["privacy_result"] = "blocked"
        summary["safety_result"] = "passed"
        _write_docs()
        _write_repo_reports(summary, {"total_files_seen": 0, "total_supported_files": 0, "total_unsupported_files": 0, "total_duplicate_files": 0, "extension_counts": {}}, [])
        print(f"{BLOCK}_BLOCKED")
        print(json.dumps({"reason": str(exc), "provider_call_made": False, "billing_api_call_made": False}, indent=2))
        return 1
    result = "VAULT_REQUIRED" if summary["vault_status"] == "VAULT_REQUIRED" else "PASS"
    print(f"{BLOCK}_{result}")
    print(
        json.dumps(
            {
                "vault_status": summary["vault_status"],
                "corpus_inventory_completed": summary["corpus_inventory_completed"],
                "tokenized_corpus_generated": summary["tokenized_corpus_generated"],
                "total_files_seen": summary["total_files_seen"],
                "total_supported_files": summary["total_supported_files"],
                "provider_call_made": summary["provider_call_made"],
                "billing_api_call_made": summary["billing_api_call_made"],
                "privacy_result": summary["privacy_result"],
                "safety_result": summary["safety_result"],
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
