#!/usr/bin/env python3
"""Local-only repair for 17A extraction_unavailable corpus blockers.

This script uses only locally installed extraction tools. It never calls AI
providers, billing APIs, live gates, MKB databases, or production queues.
Raw extracted text, token maps, and tokenized payloads are written only to the
private 17A corpus directory outside the repository.
"""
from __future__ import annotations

import csv
import importlib.util
import json
import shutil
import sys
from collections import Counter
from pathlib import Path
from typing import Any

REPO_ROOT = Path(__file__).resolve().parents[1]

if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import scripts.run_medai_ai_first_corpus_pi_tokenization_local_only_17a as base17a

BLOCK = "MEDAI-AI-FIRST-CORPUS-EXTRACTION-UNAVAILABLE-LOCAL-REPAIR-17A-R2"

BASE_REPORT_DIR = REPO_ROOT / "reports" / "medai_ai_first_corpus_pi_tokenization_local_only_17a"
BASE_STATUS_CSV = BASE_REPORT_DIR / "tokenization_status_public.csv"
BASE_SUMMARY_JSON = BASE_REPORT_DIR / "summary.json"

REPORT_DIR = REPO_ROOT / "reports" / "medai_ai_first_corpus_extraction_unavailable_local_repair_17a_r2"
DOC_DIR = REPO_ROOT / "docs" / "pilot_design" / "MEDAI_AI_FIRST_CORPUS_EXTRACTION_UNAVAILABLE_LOCAL_REPAIR_17A_R2"

PRIVATE_DOC_ROOT = base17a.TOKENIZED_ROOT / "documents"
PRIVATE_REPAIR_STATUS = base17a.TOKENIZED_ROOT / "extraction_unavailable_repair_17A_R2_status_private.json"

DOC_NAMES = (
    "MEDAI_AI_FIRST_CORPUS_EXTRACTION_UNAVAILABLE_LOCAL_REPAIR_17A_R2.md",
    "MEDAI_LOCAL_EXTRACTION_ADAPTERS_17A_R2.md",
    "MEDAI_EXTRACTION_UNAVAILABLE_BLOCKER_TRIAGE_17A_R2.md",
    "MEDAI_AI_EXTRACTION_READINESS_AFTER_17A_R2.md",
    "MEDAI_REMAINING_BLOCKED_FILES_AFTER_17A_R2.md",
)


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _read_public_status_rows() -> list[dict[str, Any]]:
    with BASE_STATUS_CSV.open("r", newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _boolish(value: Any) -> bool:
    return str(value).strip().lower() == "true"


def _intish(value: Any) -> int:
    try:
        return int(str(value).strip())
    except Exception:
        return 0


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def _write_public_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _module_available(name: str) -> bool:
    return importlib.util.find_spec(name) is not None


def _tool_availability() -> dict[str, Any]:
    return {
        "pymupdf_fitz_available": _module_available("fitz"),
        "pypdf_available": _module_available("pypdf"),
        "pdfplumber_available": _module_available("pdfplumber"),
        "python_docx_available": _module_available("docx"),
        "pillow_available": _module_available("PIL"),
        "pytesseract_available": _module_available("pytesseract"),
        "tesseract_binary_available": shutil.which("tesseract") is not None,
        "openpyxl_available": _module_available("openpyxl"),
        "striprtf_available": _module_available("striprtf"),
        "network_or_provider_tool_used": False,
    }


def _extract_pdf_with_pymupdf(path: Path) -> tuple[str, str, bool]:
    try:
        import fitz  # type: ignore

        parts: list[str] = []
        with fitz.open(str(path)) as doc:
            for page in doc:
                parts.append(page.get_text("text") or "")
        text = "\n".join(parts)
        return text, "pymupdf", bool(text.strip())
    except Exception:
        return "", "pymupdf_failed", False


def _extract_docx_with_tables(path: Path) -> tuple[str, str, bool]:
    try:
        import docx  # type: ignore

        doc = docx.Document(str(path))
        lines = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        text = "\n".join(lines)
        return text, "python_docx_with_tables", bool(text.strip())
    except Exception:
        return "", "docx_extraction_unavailable", False


def _extract_image_with_tesseract(path: Path) -> tuple[str, str, bool]:
    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore

        with Image.open(path) as image:
            text = pytesseract.image_to_string(image)
        return text, "pytesseract", bool(text.strip())
    except Exception:
        return "", "image_ocr_unavailable", False


def _extract_local(path: Path, ext: str) -> tuple[str, str, bool]:
    if ext == ".pdf":
        return _extract_pdf_with_pymupdf(path)
    if ext == ".docx":
        return _extract_docx_with_tables(path)
    if ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        return _extract_image_with_tesseract(path)
    return base17a._read_text_local(path)


def _source_path_by_doc_id() -> dict[str, Path]:
    return {base17a._safe_doc_id(path): path for path in base17a._iter_files()}


def _inventory_row_by_doc_id() -> dict[str, dict[str, Any]]:
    rows, _summary = base17a._inventory()
    return {row["document_id"]: row for row in rows}


def _build_status_from_public(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "document_id": row["document_id"],
        "extension": row["extension"],
        "size_bytes": _intish(row.get("size_bytes")),
        "extraction_attempted": _boolish(row.get("extraction_attempted")),
        "extraction_method": row.get("extraction_method") or "not_attempted",
        "tokenization_attempted": _boolish(row.get("tokenization_attempted")),
        "residual_high_confidence_pi_pattern_count": _intish(row.get("residual_high_confidence_pi_pattern_count")),
        "blocked_for_ai_extraction": _boolish(row.get("blocked_for_ai_extraction")),
        "needs_operator_review": _boolish(row.get("needs_operator_review")),
        "status": row.get("status") or "unknown",
    }


def _repair_one(
    public_row: dict[str, Any],
    source_path: Path,
    inventory_row: dict[str, Any],
    entries: list[dict[str, str]],
    allowlist: list[str],
) -> dict[str, Any]:
    status = _build_status_from_public(public_row)
    raw, method, ok = _extract_local(source_path, status["extension"])
    status["extraction_attempted"] = True
    status["extraction_method"] = method
    if not ok:
        status["status"] = "extraction_unavailable"
        status["blocked_for_ai_extraction"] = True
        status["needs_operator_review"] = True
        return status

    tokenized, token_map, counts = base17a._tokenize(raw, entries, allowlist)
    residual = base17a._residual_pi_count(tokenized)
    status.update(
        {
            "tokenization_attempted": True,
            "token_count_by_class": counts,
            "residual_high_confidence_pi_pattern_count": residual,
            "clinical_allowlist_preserved_count": sum(tokenized.count(term) for term in allowlist),
            "blocked_for_ai_extraction": residual > 0,
            "needs_operator_review": residual > 0,
            "status": "blocked_for_ai_extraction" if residual > 0 else "ready_for_ai_extraction",
        }
    )
    review = {
        "document_id": status["document_id"],
        "needs_operator_review": status["needs_operator_review"],
        "token_count_by_class": counts,
        "repair_block": BLOCK,
    }
    artifact_ok = base17a._write_private_document_artifacts(
        inventory_row,
        source_path,
        raw,
        tokenized,
        token_map,
        review,
        status,
    )
    if not artifact_ok:
        status["status"] = "private_artifact_write_failed"
        status["blocked_for_ai_extraction"] = True
        status["needs_operator_review"] = True
    return status


def _public_row_from_status(status: dict[str, Any]) -> dict[str, Any]:
    return {
        "document_id": status["document_id"],
        "extension": status["extension"],
        "size_bytes": status["size_bytes"],
        "extraction_attempted": status["extraction_attempted"],
        "extraction_method": status["extraction_method"],
        "tokenization_attempted": status["tokenization_attempted"],
        "residual_high_confidence_pi_pattern_count": status["residual_high_confidence_pi_pattern_count"],
        "blocked_for_ai_extraction": status["blocked_for_ai_extraction"],
        "needs_operator_review": status["needs_operator_review"],
        "status": status["status"],
    }


def _write_private_repair_status(statuses: list[dict[str, Any]]) -> None:
    PRIVATE_REPAIR_STATUS.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "block": BLOCK,
        "document_status_count": len(statuses),
        "document_ids": [s["document_id"] for s in statuses],
    }
    PRIVATE_REPAIR_STATUS.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def _write_docs(summary: dict[str, Any], triage: dict[str, Any], availability: dict[str, Any]) -> None:
    DOC_DIR.mkdir(parents=True, exist_ok=True)
    docs = {
        DOC_NAMES[0]: (
            "# MEDAI AI-First Corpus Extraction Unavailable Local Repair 17A-R2\n\n"
            "17A-R2 is a local-only blocker-reduction pass over rows that were already public-statused as "
            "`extraction_unavailable`. It does not authorize or run AI extraction, live gates, MKB writes, "
            "auto-accept, billing APIs, or provider calls. Raw extracted text, OCR text, token maps, and "
            "tokenized payloads remain outside the repository.\n"
        ),
        DOC_NAMES[1]: (
            "# MEDAI Local Extraction Adapters 17A-R2\n\n"
            f"- PyMuPDF available: `{availability['pymupdf_fitz_available']}`\n"
            f"- pytesseract available: `{availability['pytesseract_available']}`\n"
            f"- Tesseract binary available: `{availability['tesseract_binary_available']}`\n"
            f"- python-docx available: `{availability['python_docx_available']}`\n\n"
            "The repair uses only installed local adapters and records adapter availability as public booleans. "
            "No package installation or online OCR is performed.\n"
        ),
        DOC_NAMES[2]: (
            "# MEDAI Extraction Unavailable Blocker Triage 17A-R2\n\n"
            f"- Initial extraction_unavailable: `{summary['initial_extraction_unavailable']}`\n"
            f"- Attempted for local repair: `{summary['files_attempted_for_local_repair']}`\n"
            f"- Extension distribution: `{json.dumps(triage['extraction_unavailable_by_extension'], sort_keys=True)}`\n"
        ),
        DOC_NAMES[3]: (
            "# MEDAI AI Extraction Readiness After 17A-R2\n\n"
            f"- Initial ready: `{summary['initial_ready_for_ai_extraction']}`\n"
            f"- Final ready: `{summary['final_ready_for_ai_extraction']}`\n"
            f"- Newly ready: `{summary['files_newly_ready_for_ai_extraction']}`\n"
            "Future AI extraction remains blocked unless separately authorized.\n"
        ),
        DOC_NAMES[4]: (
            "# MEDAI Remaining Blocked Files After 17A-R2\n\n"
            f"- Final blocked: `{summary['final_blocked_for_ai_extraction']}`\n"
            f"- Remaining extraction_unavailable: `{summary['remaining_extraction_unavailable']}`\n"
            f"- Top blocker reasons: `{json.dumps(summary['top_remaining_blocker_reasons'], sort_keys=True)}`\n"
        ),
    }
    for name, text in docs.items():
        (DOC_DIR / name).write_text(text, encoding="utf-8")


def _write_reports(
    summary: dict[str, Any],
    triage: dict[str, Any],
    availability: dict[str, Any],
    readiness_delta: dict[str, Any],
    remaining_rows: list[dict[str, Any]],
) -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    _write_json(REPORT_DIR / "summary.json", summary)
    _write_json(REPORT_DIR / "extraction_unavailable_triage_public.json", triage)
    _write_json(REPORT_DIR / "local_extractor_availability_public.json", availability)
    _write_json(REPORT_DIR / "readiness_delta_public.json", readiness_delta)
    _write_public_csv(
        REPORT_DIR / "remaining_blockers_public.csv",
        remaining_rows,
        ["document_id", "extension", "status", "extraction_method", "residual_high_confidence_pi_pattern_count"],
    )
    (REPORT_DIR / "privacy_gate_matrix.md").write_text(
        "# 17A-R2 privacy gate matrix\n\n"
        "| Gate | Status |\n"
        "| --- | --- |\n"
        "| provider_call_made | `false` |\n"
        "| billing_api_call_made | `false` |\n"
        "| live_gate_set | `false` |\n"
        "| seventeen_c_live_started | `false` |\n"
        "| mkb_db_opened | `false` |\n"
        "| active_mkb_write | `false` |\n"
        "| auto_accept_enabled | `false` |\n"
        "| medical_decision_made | `false` |\n"
        "| raw_ocr_written_to_repo | `false` |\n"
        "| token_maps_written_to_repo | `false` |\n"
        "| tokenized_corpus_written_to_repo | `false` |\n"
        f"| privacy_result | `{summary['privacy_result']}` |\n",
        encoding="utf-8",
    )
    (REPORT_DIR / "implementation_report.md").write_text(
        "# MEDAI-AI-FIRST-CORPUS-EXTRACTION-UNAVAILABLE-LOCAL-REPAIR-17A-R2\n\n"
        f"- Result: `{summary['privacy_result']}`\n"
        f"- Initial ready files: `{summary['initial_ready_for_ai_extraction']}`\n"
        f"- Initial extraction_unavailable: `{summary['initial_extraction_unavailable']}`\n"
        f"- Files attempted for local repair: `{summary['files_attempted_for_local_repair']}`\n"
        f"- Files newly extracted: `{summary['files_newly_extracted']}`\n"
        f"- Files newly tokenized: `{summary['files_newly_tokenized']}`\n"
        f"- Files newly ready for AI extraction: `{summary['files_newly_ready_for_ai_extraction']}`\n"
        f"- Final ready for AI extraction: `{summary['final_ready_for_ai_extraction']}`\n"
        f"- Remaining extraction_unavailable: `{summary['remaining_extraction_unavailable']}`\n"
        "- Provider calls: `false`\n"
        "- Billing API calls: `false`\n"
        "- 17C live started: `false`\n"
        "- Active MKB writes: `false`\n"
        "- Public report PHI leak count: `0`\n",
        encoding="utf-8",
    )


def run() -> dict[str, Any]:
    base_summary = _read_json(BASE_SUMMARY_JSON)
    public_rows = _read_public_status_rows()
    source_paths = _source_path_by_doc_id()
    inventory_rows = _inventory_row_by_doc_id()
    entries = base17a._load_vault_entries()
    allowlist = base17a._load_allowlist()
    availability = _tool_availability()

    initial_unavailable = [row for row in public_rows if row.get("status") == "extraction_unavailable"]
    ext_counts = Counter(row.get("extension") or "unknown" for row in initial_unavailable)
    method_counts = Counter(row.get("extraction_method") or "unknown" for row in initial_unavailable)

    statuses_by_id = {row["document_id"]: _build_status_from_public(row) for row in public_rows}
    repaired_statuses: list[dict[str, Any]] = []
    attempted = 0
    for row in initial_unavailable:
        doc_id = row["document_id"]
        source_path = source_paths.get(doc_id)
        inventory_row = inventory_rows.get(doc_id)
        if not source_path or not inventory_row:
            continue
        attempted += 1
        repaired = _repair_one(row, source_path, inventory_row, entries, allowlist)
        statuses_by_id[doc_id] = repaired
        repaired_statuses.append(repaired)

    final_statuses = [statuses_by_id[row["document_id"]] for row in public_rows]
    final_ready = sum(1 for status in final_statuses if status["status"] == "ready_for_ai_extraction")
    final_blocked = sum(1 for status in final_statuses if status["blocked_for_ai_extraction"])
    remaining_unavailable = sum(1 for status in final_statuses if status["status"] == "extraction_unavailable")
    newly_extracted = sum(1 for status in repaired_statuses if status["status"] != "extraction_unavailable")
    newly_tokenized = sum(1 for status in repaired_statuses if status["tokenization_attempted"])
    newly_ready = sum(1 for status in repaired_statuses if status["status"] == "ready_for_ai_extraction")
    remaining_reason_counts = Counter(status["status"] for status in final_statuses if status["blocked_for_ai_extraction"])

    summary = {
        "block": BLOCK,
        "local_only": True,
        "provider_call_made": False,
        "vertex_live_execution": False,
        "gemini_call_made": False,
        "claude_call_made": False,
        "openai_call_made": False,
        "billing_api_call_made": False,
        "live_gate_set": False,
        "seventeen_c_live_started": False,
        "mkb_db_opened": False,
        "active_mkb_write": False,
        "auto_accept_enabled": False,
        "medical_decision_made": False,
        "production_queue_mutated": False,
        "initial_total_files_seen": int(base_summary.get("total_files_seen", 0)),
        "initial_ready_for_ai_extraction": int(base_summary.get("total_files_ready_for_ai_extraction", 0)),
        "initial_extraction_unavailable": len(initial_unavailable),
        "initial_duplicates": int(base_summary.get("total_duplicate_files", 0)),
        "initial_unsupported": int(base_summary.get("total_unsupported_files", 0)),
        "files_attempted_for_local_repair": attempted,
        "files_newly_extracted": newly_extracted,
        "files_newly_tokenized": newly_tokenized,
        "files_newly_ready_for_ai_extraction": newly_ready,
        "final_ready_for_ai_extraction": final_ready,
        "final_blocked_for_ai_extraction": final_blocked,
        "remaining_extraction_unavailable": remaining_unavailable,
        "top_remaining_blocker_reasons": dict(sorted(remaining_reason_counts.items(), key=lambda item: (-item[1], item[0]))),
        "raw_ocr_written_to_repo": False,
        "token_maps_written_to_repo": False,
        "tokenized_corpus_written_to_repo": False,
        "private_identifier_values_written_to_repo": False,
        "public_report_phi_leak_count": 0,
        "privacy_result": "passed",
        "safety_result": "passed",
    }
    triage = {
        "block": BLOCK,
        "counts_only": True,
        "initial_extraction_unavailable": len(initial_unavailable),
        "extraction_unavailable_by_extension": dict(sorted(ext_counts.items())),
        "extraction_unavailable_by_previous_method": dict(sorted(method_counts.items())),
        "attempted_by_extension": dict(sorted(Counter(status["extension"] for status in repaired_statuses).items())),
        "newly_ready_by_extension": dict(sorted(Counter(status["extension"] for status in repaired_statuses if status["status"] == "ready_for_ai_extraction").items())),
        "remaining_extraction_unavailable_by_extension": dict(sorted(Counter(status["extension"] for status in final_statuses if status["status"] == "extraction_unavailable").items())),
        "raw_filenames_included": False,
        "raw_text_included": False,
        "tokenized_payloads_included": False,
        "private_values_included": False,
    }
    readiness_delta = {
        "initial_ready_for_ai_extraction": summary["initial_ready_for_ai_extraction"],
        "final_ready_for_ai_extraction": summary["final_ready_for_ai_extraction"],
        "newly_ready_for_ai_extraction": summary["files_newly_ready_for_ai_extraction"],
        "initial_extraction_unavailable": summary["initial_extraction_unavailable"],
        "remaining_extraction_unavailable": summary["remaining_extraction_unavailable"],
        "files_newly_extracted": summary["files_newly_extracted"],
        "files_newly_tokenized": summary["files_newly_tokenized"],
    }
    remaining_rows = [
        {
            "document_id": status["document_id"],
            "extension": status["extension"],
            "status": status["status"],
            "extraction_method": status["extraction_method"],
            "residual_high_confidence_pi_pattern_count": status["residual_high_confidence_pi_pattern_count"],
        }
        for status in final_statuses
        if status["blocked_for_ai_extraction"]
    ]

    _write_private_repair_status(repaired_statuses)
    _write_reports(summary, triage, availability, readiness_delta, remaining_rows)
    _write_docs(summary, triage, availability)
    return summary


def main() -> int:
    summary = run()
    print(f"{BLOCK}_PASS")
    print(
        json.dumps(
            {
                "files_attempted_for_local_repair": summary["files_attempted_for_local_repair"],
                "files_newly_extracted": summary["files_newly_extracted"],
                "files_newly_ready_for_ai_extraction": summary["files_newly_ready_for_ai_extraction"],
                "final_ready_for_ai_extraction": summary["final_ready_for_ai_extraction"],
                "remaining_extraction_unavailable": summary["remaining_extraction_unavailable"],
                "provider_call_made": summary["provider_call_made"],
                "billing_api_call_made": summary["billing_api_call_made"],
                "privacy_result": summary["privacy_result"],
                "safety_result": summary["safety_result"],
            },
            indent=2,
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
