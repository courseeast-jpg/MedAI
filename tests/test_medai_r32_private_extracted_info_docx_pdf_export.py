"""Tests for R32 private extracted-info DOCX/PDF export.

Validates the counts-only public report and the locally-generated (uncommitted) DOCX.
Asserts the private export is not tracked by git and the public report has no clinical text.
"""
from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = REPO_ROOT / "reports" / "medai_r32_private_extracted_info_docx_pdf_export"
DOCX_PATH = REPO_ROOT / "private_exports" / "medai_extracted_info_r32" / "MedAI_Extracted_Info_R32.docx"
PUBLIC_REPORTS = ("summary.json", "export_counts_public.json", "implementation_report.md")
SECRET_PATTERNS = [
    re.compile(r"AIza[0-9A-Za-z_-]{20,}"), re.compile(r"ya29\."), re.compile(r"Bearer "),
    re.compile(r"[A-Za-z]:\\\\Users\\\\"), re.compile(r"\bMRN\b"), re.compile(r"\bDOB\b"),
]


def _summary() -> dict:
    return json.loads((REPORT_DIR / "summary.json").read_text(encoding="utf-8"))


def test_public_reports_exist():
    for n in PUBLIC_REPORTS:
        assert (REPORT_DIR / n).exists(), n


def test_overall_pass_and_counts():
    s = _summary()
    assert s["overall_result"] == "PASS"
    assert s["docx_created"] is True
    assert s["extracted_payload_records_exported"] == 179
    assert s["not_extracted_records_indexed"] == 317
    assert s["total_staging_records"] == 496
    assert s["usable_payload_records"] + s["empty_extraction_shell_records"] == 179


def test_docx_path_and_pdf_reported():
    s = _summary()
    assert s["docx_path"] == "private_exports/medai_extracted_info_r32/MedAI_Extracted_Info_R32.docx"
    # PDF optional: either created with a path, or reported unavailable.
    assert (s["pdf_created"] is True) or (s["pdf_path"] == "unavailable")


def test_docx_content_structure():
    if not DOCX_PATH.is_file():
        pytest.skip("DOCX not present (run the export script first)")
    from docx import Document
    text = "\n".join(p.text for p in Document(str(DOCX_PATH)).paragraphs)
    assert "MedAI Extracted Information Export — R32" in text          # title
    assert "Extracted payload records: 179" in text                   # summary count
    assert "Extracted payload records (179)" in text                  # payload section
    assert "Extracted records with no usable items" in text           # empty-shell section
    assert "Not-extracted records index (317)" in text                # appendix count
    assert sum(1 for p in Document(str(DOCX_PATH)).paragraphs if p.text.startswith("Record ")) >= 1
    s = _summary()
    assert s["full_schema_record_present"] is True
    if s["minimal_review_record_present"]:
        assert s["minimal_review_record_present"] is True


def test_private_export_not_tracked_by_git():
    s = _summary()
    assert s["private_export_committed"] is False
    out = subprocess.run(["git", "ls-files", "private_exports/"], cwd=REPO_ROOT,
                         capture_output=True, text=True, timeout=30)
    assert out.stdout.strip() == "", "private export must not be tracked by git"


def test_safety_flags():
    s = _summary()
    for k in ("provider_model_call_made", "live_extraction_started", "new_extraction_started",
              "auto_accept_enabled", "medical_decision_made"):
        assert s[k] is False, k
    assert s["active_verified_records_created"] == 0


def test_public_reports_counts_only_no_clinical_text():
    s = _summary()
    assert s["raw_clinical_text_in_public_report"] is False
    assert s["public_report_phi_leak_count"] == 0
    assert s["private_path_leaks_after"] == 0
    assert s["secret_leaks_after"] == 0
    for n in PUBLIC_REPORTS:
        text = (REPORT_DIR / n).read_text(encoding="utf-8", errors="ignore")
        for pat in SECRET_PATTERNS:
            assert not pat.search(text), (n, pat.pattern)
        # Public reports must not dump per-record extracted content.
        assert "Extracted items / facts" not in text
        assert not re.search(r"\b\d{3}[- ]\d{2}[- ]\d{4}\b", text)
